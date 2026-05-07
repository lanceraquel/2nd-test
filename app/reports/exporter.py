import json
from io import BytesIO
from pathlib import Path

from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import PipelineStage, ReportArtifact, ResearchTask


def _report_dir(task_id: int) -> Path:
    settings = get_settings()
    directory = Path(settings.report_output_dir) / f"task-{task_id}"
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def export_task_reports(session: Session, task: ResearchTask) -> dict[str, str]:
    stages = list(
        session.scalars(
            select(PipelineStage).where(PipelineStage.task_id == task.id).order_by(PipelineStage.stage_order.asc(), PipelineStage.attempt.asc())
        ).all()
    )
    directory = _report_dir(task.id)
    paths = {
        "json": str(directory / "run.json"),
        "markdown": str(directory / "report.md"),
        "html": str(directory / "report.html"),
    }
    payload = {
        "task": {
            "id": task.id,
            "title": task.title,
            "brief": task.brief,
            "research_questions": task.research_questions,
            "entities": task.entities,
            "focus_areas": task.focus_areas,
            "status": task.status,
            "current_stage": task.current_stage,
        },
        "stages": [
            {
                "stage_key": stage.stage_key,
                "stage_title": stage.stage_title,
                "phase": stage.phase,
                "stage_order": stage.stage_order,
                "attempt": stage.attempt,
                "status": stage.status,
                "summary": stage.summary,
                "payload_json": stage.payload_json,
            }
            for stage in stages
        ],
        "final_report_markdown": task.final_report_markdown,
        "final_report_html": task.final_report_html,
    }
    Path(paths["json"]).write_text(json.dumps(payload, indent=2), encoding="utf-8")
    Path(paths["markdown"]).write_text(task.final_report_markdown or "# Report pending approval", encoding="utf-8")
    Path(paths["html"]).write_text(task.final_report_html or "<article><p>HTML report pending approval.</p></article>", encoding="utf-8")
    _store_docx(session, task)
    return paths


def _store_docx(session: Session, task: ResearchTask) -> None:
    document = Document()
    document.add_heading(task.title, level=0)
    document.add_paragraph(task.brief)
    document.add_heading("Research Questions", level=1)
    for question in task.research_questions or ["No explicit research questions supplied."]:
        document.add_paragraph(question, style="List Bullet")

    document.add_heading("Final Report", level=1)
    for block in (task.final_report_markdown or "Report pending approval.").split("\n\n"):
        if block.strip():
            document.add_paragraph(block.strip())

    buffer = BytesIO()
    document.save(buffer)
    artifact = ReportArtifact(
        task_id=task.id,
        format="docx",
        filename=f"atlas-research-run-{task.id}.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=buffer.getvalue(),
    )
    session.add(artifact)
    session.commit()
